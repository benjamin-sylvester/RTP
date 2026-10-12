"""Import 6.0 Claude/capital/RTP_Relationship_Tracker.xlsx into organizations, people, people_emails,
interactions (one per 'Last touch'), action_items (Next actions tab + per-row Next step) and
investments (LPs tab '186 John' column -> 186 John Street LLC).

Idempotent: people match on email, then on lower(full_name); interactions key on
(source='tracker:<tab>', source_id=row); action items key on (source, source_id).
  python scripts/import_tracker.py [--sandbox] [--commit]
"""
import datetime as dt
import os
import re
import sys

import openpyxl
import psycopg
from _conn import connect, load_env, ROOT

XLSX = ROOT.parent / "capital" / "RTP_Relationship_Tracker.xlsx"
TODAY = dt.date(2026, 10, 12)


def conn_for(argv):
    if "--sandbox" in argv:
        load_env()
        return psycopg.connect(os.environ["SANDBOX_DATABASE_URL"], autocommit=False)
    return connect(autocommit=False)


def emails_in(s):
    return [e.lower() for e in re.findall(r"[\w.+-]+@[\w-]+\.[\w.-]+", s or "")]


def to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    m = re.fullmatch(r"(\d{4})-(\d{1,2})(?:-(\d{1,2}))?", s)
    if m:
        return dt.date(int(m[1]), int(m[2]), int(m[3] or 1))
    m = re.search(r"(\d{1,2})/(\d{1,2})(?:/(\d{2,4}))?", s)
    if m:
        y = int(m[3]) if m[3] else 2026
        y = y + 2000 if y < 100 else y
        try:
            return dt.date(y, int(m[1]), int(m[2]))
        except ValueError:
            return None
    if s.lower().startswith("now"):
        return TODAY
    return None


def cents(v):
    if v is None or v == "":
        return None
    try:
        return int(round(float(str(v).replace("$", "").replace(",", "")) * 100))
    except ValueError:
        return None


class Importer:
    def __init__(self, c, log=print):
        self.c, self.log = c, log
        self.n = {"org": 0, "person": 0, "email": 0, "interaction": 0, "action": 0, "investment": 0, "matched": 0}
        self.ben = c.execute("SELECT id FROM people WHERE full_name='Ben Sylvester'").fetchone()[0]
        self.john = c.execute("SELECT id FROM entities WHERE short_name='John St LLC'").fetchone()[0]

    def org(self, name, kind="other", source=None):
        name = (name or "").strip().rstrip(".,;")
        if not name:
            return None
        r = self.c.execute("SELECT id FROM organizations WHERE lower(name)=lower(%s) ORDER BY id LIMIT 1", (name,)).fetchone()
        if r:
            return r[0]
        self.n["org"] += 1
        return self.c.execute("INSERT INTO organizations (name, kind, source) VALUES (%s,%s,%s) RETURNING id",
                              (name, kind, source)).fetchone()[0]

    def person(self, full_name, emails=(), roles=(), org_id=None, source=None, source_id=None, **fields):
        full_name = re.sub(r"\s+", " ", (full_name or "")).strip()
        if not full_name:
            return None
        pid = None
        for e in emails:
            r = self.c.execute("SELECT person_id FROM people_emails WHERE email=%s", (e,)).fetchone()
            if r:
                pid = r[0]; break
        if pid is None:
            r = self.c.execute("SELECT id FROM people WHERE lower(full_name)=lower(%s) AND merged_into_id IS NULL ORDER BY id LIMIT 1",
                               (full_name,)).fetchone()
            pid = r[0] if r else None
        if pid is None:
            parts = full_name.split(" ")
            pid = self.c.execute(
                "INSERT INTO people (full_name, first_name, last_name, org_id, roles, source, source_id) "
                "VALUES (%s,%s,%s,%s,%s,%s,%s) RETURNING id",
                (full_name, parts[0], parts[-1] if len(parts) > 1 else None, org_id, list(roles), source, source_id)).fetchone()[0]
            self.n["person"] += 1
        else:
            self.n["matched"] += 1
            if roles:
                self.c.execute("UPDATE people SET roles = (SELECT array_agg(DISTINCT r) FROM unnest(roles || %s::text[]) r) WHERE id=%s",
                               (list(roles), pid))
            if org_id:
                self.c.execute("UPDATE people SET org_id = COALESCE(org_id, %s) WHERE id=%s", (org_id, pid))
        for k, v in fields.items():
            if v not in (None, ""):
                self.c.execute(f"UPDATE people SET {k} = %s WHERE id=%s", (v, pid))
        for i, e in enumerate(emails):
            if not self.c.execute("SELECT 1 FROM people_emails WHERE email=%s", (e,)).fetchone():
                self.c.execute("INSERT INTO people_emails (person_id, email, is_primary) VALUES (%s,%s,%s)", (pid, e, i == 0))
                self.n["email"] += 1
        return pid

    def note(self, pid, text):
        if pid and text:
            self.c.execute("UPDATE people SET notes = CASE WHEN notes IS NULL THEN %s WHEN position(%s in notes) > 0 THEN notes ELSE notes || E'\\n' || %s END WHERE id=%s",
                           (text, text[:80], text, pid))

    def interaction(self, pid, when, summary, source, source_id, channel="other", org_id=None):
        when = to_date(when)
        if not (pid or org_id) or not when or not summary:
            return
        if self.c.execute("SELECT 1 FROM interactions WHERE source=%s AND source_id=%s AND person_id IS NOT DISTINCT FROM %s",
                          (source, source_id, pid)).fetchone():
            return
        self.c.execute("INSERT INTO interactions (person_id, org_id, channel, direction, occurred_at, summary, by_person_id, source, source_id) "
                       "VALUES (%s,%s,%s,'mutual',%s,%s,%s,%s,%s)",
                       (pid, org_id, channel, dt.datetime.combine(when, dt.time(12)), summary[:1000], self.ben, source, source_id))
        self.n["interaction"] += 1

    def action(self, title, pid, due, status, details, source, source_id, org_id=None):
        title = (title or "").strip()
        if not title:
            return
        if self.c.execute("SELECT 1 FROM action_items WHERE source=%s AND source_id=%s", (source, source_id)).fetchone():
            return
        self.c.execute("INSERT INTO action_items (title, details, owner_person_id, due_on, status, person_id, org_id, source, source_id, completed_at) "
                       "VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)",
                       (title[:300], details, self.ben, to_date(due), status, pid, org_id, source, source_id,
                        dt.datetime.now() if status == "done" else None))
        self.n["action"] += 1

    def investment(self, pid, funded_cents, notes, source, source_id):
        if self.c.execute("SELECT 1 FROM investments WHERE entity_id=%s AND person_id=%s", (self.john, pid)).fetchone():
            return
        self.c.execute("INSERT INTO investments (entity_id, person_id, funded_cents, notes, source, source_id) VALUES (%s,%s,%s,%s,%s,%s)",
                       (self.john, pid, funded_cents, notes, source, source_id))
        self.n["investment"] += 1


def rows(ws, header_row):
    hdr = [str(c).strip() if c is not None else "" for c in next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))]
    for i, r in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        d = dict(zip(hdr, r))
        if any(v not in (None, "") for v in r):
            yield i, d


def hyperlinks(ws, header_row, col_names):
    """Map row -> {col_name: url} for hyperlink cells (needs a non-read-only workbook)."""
    hdr = [str(c.value).strip() if c.value is not None else "" for c in ws[header_row]]
    idx = {n: hdr.index(n) for n in col_names if n in hdr}
    out = {}
    for row in ws.iter_rows(min_row=header_row + 1):
        for n, j in idx.items():
            cell = row[j]
            if cell.hyperlink and cell.hyperlink.target:
                out.setdefault(cell.row, {})[n] = cell.hyperlink.target
    return out


def main():
    c = conn_for(sys.argv)
    im = Importer(c)
    wb = openpyxl.load_workbook(XLSX, data_only=True)

    # Capital partners
    ws = wb["Capital partners"]
    for i, d in rows(ws, 4):
        name = d.get("Name") or ""
        if str(d.get("Priority")).strip().lower() == "merged":
            continue
        src, sid = "tracker:Capital partners", str(i)
        m = re.match(r"(.+?)\s*\((.+)\)\s*$", name)
        emails = emails_in(d.get("Email"))
        wraw = (d.get("Warmth") or "").strip().lower()
        warmth = wraw.split()[0] if wraw and wraw.split()[0] in ("hot", "warm", "cold", "closed") else None
        tier = d.get("Priority") if d.get("Priority") in ("A", "B", "C", "D") else None
        notes = " | ".join(str(d[k]) for k in ("Status", "Notes") if d.get(k))
        if wraw and warmth is None:
            notes = f"warmth: {wraw} | " + notes
        if m:  # "Firm (Person, Person)"
            oid = im.org(m[1], "capital_partner", src)
            people_names = [p.strip() for p in m[2].split(",")]
            def mail_for(pn, k):
                first = pn.split()[0].lower()
                hits = [e for e in emails if first[:4] in e.split("@")[0].lower()]
                if hits:
                    return hits[:1]
                unclaimed = [e for e in emails if not any(q.split()[0].lower()[:4] in e.split("@")[0].lower() for q in people_names)]
                return unclaimed[k:k + 1] if k < len(unclaimed) else ()
            for k, pn in enumerate(people_names):
                pid = im.person(pn, mail_for(pn, k), ("capital_partner",), oid, src, sid,
                                warmth=warmth, lp_tier=tier)
                im.note(pid, notes)
                if k == 0:
                    im.interaction(pid, d.get("Last touch"), notes or f"Last touch ({name})", src, sid, org_id=oid)
                    im.action(d.get("Next step"), pid, None, "open", notes, src, sid, org_id=oid)
        else:
            firm = (d.get("Type / firm") or "").split("/")[0].split(",")[0].strip()
            looks_like_person = len(name.split()) in (2, 3) and not re.search(r"(LLC|Capital|Partners|Co\.|Company|Group)", name)
            oid = im.org(firm if looks_like_person else name, "capital_partner", src) if (firm or not looks_like_person) else None
            if looks_like_person:
                pid = im.person(name, emails, ("capital_partner",), oid, src, sid, warmth=warmth, lp_tier=tier)
                im.note(pid, notes)
                im.interaction(pid, d.get("Last touch"), notes or "Last touch", src, sid, org_id=oid)
                im.action(d.get("Next step"), pid, None, "open", notes, src, sid, org_id=oid)
            else:
                im.c.execute("UPDATE organizations SET notes=%s WHERE id=%s", (notes, oid))
                im.interaction(None, d.get("Last touch"), notes or "Last touch", src, sid, org_id=oid)
                im.action(d.get("Next step"), None, None, "open", notes, src, sid, org_id=oid)

    # LPs
    ws = wb["LPs"]
    for i, d in rows(ws, 4):
        src, sid = "tracker:LPs", str(i)
        name = d.get("Name") or ""
        emails = emails_in(d.get("Email"))
        invested = str(d.get("186 John") or "").strip().lower().startswith("invest")
        tier = d.get("Tier") if d.get("Tier") in ("A", "B", "C", "D") else None
        roles = ("lp",) if invested else ("lp_prospect",)
        pid = im.person(name, emails, roles, None, src, sid, lp_tier=tier,
                        lp_relationship=d.get("Relationship"), lp_indicated_cents=cents(d.get("Indicated $")),
                        lp_next_deal_stage=d.get("Next deal stage"),
                        warmth="closed" if tier == "D" else None)
        im.note(pid, f"[{d.get('Category')}] {d.get('Notes') or ''}".strip())
        im.interaction(pid, d.get("Last touch"), d.get("Notes") or "Last touch", src, sid)
        im.action(d.get("Next step"), pid, None, "open", d.get("Notes"), src, sid)
        if invested:
            im.investment(pid, cents(d.get("186 $ in")), f"From tracker LPs tab; '186 $ in' = {d.get('186 $ in')}", src, sid)

    # Reconvene
    ws = wb["Reconvene"]
    links = hyperlinks(ws, 4, ["Name", "LinkedIn"])
    for i, d in rows(ws, 4):
        src, sid = "tracker:Reconvene", str(i)
        typ = (d.get("Type") or "").lower()
        roles = ("lp_prospect",) if "lp" in typ else ("connector",) if "connect" in typ else ("operator_partner",) if ("operator" in typ or "gp" in typ) else ()
        oid = im.org(re.sub(r"\s*\(.*\)$", "", d.get("Firm") or ""), "capital_partner" if "lp" in typ else "other", src)
        pid = im.person(d.get("Name"), (), roles, oid, src, sid,
                        reconvene_url=links.get(i, {}).get("Name"), linkedin_url=links.get(i, {}).get("LinkedIn"))
        im.note(pid, f"[Reconvene 2026 tier {d.get('Tier')}, {d.get('Type')}] fit: {d.get('Fit for RTP') or ''} | status: {d.get('Status') or ''} | {d.get('Notes') or ''} | met: {d.get('Met at event?') or ''}".strip())
        if str(d.get("Met at event?") or "").lower().startswith("yes") or "met" in str(d.get("Status") or "").lower():
            im.interaction(pid, "10/7/2026", f"Met at Reconvene 2026: {d.get('Notes') or d.get('Status') or ''}"[:500], src, sid, channel="in_person")
        im.action(d.get("Follow-up"), pid, None, "open", d.get("Notes"), src, sid)

    # Networking & operators
    ws = wb["Networking & operators"]
    links = hyperlinks(ws, 4, ["LinkedIn"])
    for i, d in rows(ws, 4):
        src, sid = "tracker:Networking", str(i)
        bucket = (d.get("Bucket") or "").lower()
        roles = ("operator_partner",) if "operator" in bucket else ("connector",) if "connect" in bucket else ("personal",)
        oid = im.org(re.sub(r"\s*\(.*\)$", "", (d.get("Firm") or "").split(";")[0]), "other", src)
        emails = emails_in(d.get("Contact"))
        pid = im.person(d.get("Name"), emails, roles, oid, src, sid, city=(d.get("Base") or "").split("(")[0].split("/")[0].strip() or None,
                        linkedin_url=links.get(i, {}).get("LinkedIn"))
        im.note(pid, f"[{d.get('Bucket')}] {d.get('Why / notes') or ''} | {d.get('More notes') or ''}".strip(" |"))
        im.interaction(pid, d.get("Last touch"), f"{d.get('Status') or ''}: {d.get('More notes') or d.get('Why / notes') or ''}"[:500], src, sid)
        im.action(d.get("Next step"), pid, None, "open", d.get("More notes"), src, sid)

    # Next actions (the open-items list)
    ws = wb["Next actions"]
    for i, d in rows(ws, 4):
        src, sid = "tracker:Next actions", str(i)
        who = d.get("Who") or ""
        r = c.execute("SELECT id FROM people WHERE lower(full_name)=lower(%s) AND merged_into_id IS NULL LIMIT 1", (who.strip(),)).fetchone()
        done = str(d.get("Done?") or "").strip().lower()
        status = "done" if done.startswith("done") or done in ("yes", "y", "x") else "waiting" if done.startswith("wait") else "open"
        details = " | ".join(str(d[k]) for k in ("Context", "Notes") if d.get(k))
        im.action(f"{d.get('Next step')}" + ("" if r else f" ({who})"), r[0] if r else None, d.get("When"), status,
                  f"[{d.get('Bucket')}] when: {d.get('When')} | {details}", src, sid)

    print("counts:", im.n)
    for r in c.execute("SELECT unnest(roles) AS role, count(*) FROM people GROUP BY 1 ORDER BY 2 DESC"):
        print("  role", r)
    print("  people total:", c.execute("SELECT count(*) FROM people").fetchone()[0],
          "| orgs:", c.execute("SELECT count(*) FROM organizations").fetchone()[0],
          "| emails:", c.execute("SELECT count(*) FROM people_emails").fetchone()[0],
          "| interactions:", c.execute("SELECT count(*) FROM interactions").fetchone()[0],
          "| actions open:", c.execute("SELECT count(*) FROM action_items WHERE status IN ('open','waiting')").fetchone()[0],
          "| John St investments:", c.execute("SELECT count(*), sum(funded_cents)/100 FROM investments").fetchone())
    if "--commit" in sys.argv:
        c.commit(); print("COMMITTED")
    else:
        c.rollback(); print("DRY-RUN, rolled back")


if __name__ == "__main__":
    main()
