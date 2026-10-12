"""Regenerate the relationship tracker workbook from the database (the database is the source of
truth since Phase 1; this file is a read-only view for Ben's phone).
Writes 6.0 Claude/capital/RTP_Relationship_Tracker.xlsx (archiving the previous copy first) and
logs a job_runs row ('tracker_export').
  python scripts/export_tracker.py [--sandbox] [--out <path>]
"""
import datetime as dt
import os
import shutil
import sys

import psycopg
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from _conn import connect, load_env, ROOT

OUT = ROOT.parent / "capital" / "RTP_Relationship_Tracker.xlsx"
FONT = Font(name="Times New Roman", size=11, color="000000")
HDR = Font(name="Times New Roman", size=11, bold=True, color="0000FF")
TITLE = Font(name="Times New Roman", size=11, bold=True)

TABS = {
    "Next actions": ("""
        SELECT COALESCE(to_char(due_on,'MM/DD/YY'),'') AS "When", COALESCE(about_person, about_org, about_listing, about_package, about_property, about_entity, '') AS "Who",
               title AS "Next step", COALESCE(details,'') AS "Context", owner AS "Owner", status AS "Status", to_char(updated_at,'MM/DD/YY') AS "Updated"
        FROM v_action_items_open ORDER BY due_on NULLS LAST, id""", [10, 28, 50, 60, 12, 10, 10]),
    "LPs": ("""
        SELECT COALESCE(lp_tier,'') AS "Tier", full_name AS "Name", COALESCE(organization,'') AS "Firm", COALESCE(lp_relationship,'') AS "Relationship",
               COALESCE(email,'') AS "Email", deals_invested AS "Deals in", COALESCE(funded_cents/100,0) AS "$ invested",
               COALESCE(lp_indicated_cents/100,0) AS "Indicated $", COALESCE(lp_next_deal_stage,'') AS "Next deal stage",
               COALESCE(to_char(last_interaction,'MM/DD/YY'),'') AS "Last touch", COALESCE(last_summary,'') AS "Last touch note",
               COALESCE(next_action,'') AS "Next step", COALESCE(to_char(next_action_due,'MM/DD/YY'),'') AS "Due"
        FROM v_lp_pipeline WHERE roles && ARRAY['lp','lp_prospect']::text[] ORDER BY lp_tier NULLS LAST, full_name""",
        [6, 26, 26, 22, 30, 8, 12, 12, 16, 10, 50, 40, 10]),
    "Capital partners": ("""
        SELECT COALESCE(p.lp_tier,'') AS "Priority", p.full_name AS "Name", COALESCE(o.name,'') AS "Firm", COALESCE(p.warmth,'') AS "Warmth",
               COALESCE((SELECT string_agg(email, '; ') FROM people_emails e WHERE e.person_id=p.id),'') AS "Email",
               COALESCE(to_char((SELECT max(occurred_at) FROM interactions i WHERE i.person_id=p.id),'MM/DD/YY'),'') AS "Last touch",
               COALESCE((SELECT summary FROM interactions i WHERE i.person_id=p.id ORDER BY occurred_at DESC LIMIT 1),'') AS "Last touch note",
               COALESCE((SELECT title FROM action_items a WHERE a.person_id=p.id AND a.status IN ('open','waiting') ORDER BY due_on NULLS LAST LIMIT 1),'') AS "Next step",
               COALESCE(p.notes,'') AS "Notes"
        FROM people p LEFT JOIN organizations o ON o.id=p.org_id
        WHERE 'capital_partner' = ANY(p.roles) AND p.merged_into_id IS NULL ORDER BY p.lp_tier NULLS LAST, p.full_name""",
        [8, 26, 30, 8, 34, 10, 50, 40, 60]),
    "Networking & operators": ("""
        SELECT array_to_string(array_remove(array_remove(p.roles,'lp_prospect'),'capital_partner'), ', ') AS "Roles", p.full_name AS "Name", COALESCE(o.name,'') AS "Firm",
               COALESCE(p.city,'') AS "Base", COALESCE((SELECT string_agg(email, '; ') FROM people_emails e WHERE e.person_id=p.id),'') AS "Email",
               COALESCE(p.linkedin_url,'') AS "LinkedIn",
               COALESCE(to_char((SELECT max(occurred_at) FROM interactions i WHERE i.person_id=p.id),'MM/DD/YY'),'') AS "Last touch",
               COALESCE((SELECT title FROM action_items a WHERE a.person_id=p.id AND a.status IN ('open','waiting') ORDER BY due_on NULLS LAST LIMIT 1),'') AS "Next step",
               COALESCE(p.notes,'') AS "Notes"
        FROM people p LEFT JOIN organizations o ON o.id=p.org_id
        WHERE p.roles && ARRAY['operator_partner','connector','personal']::text[] AND p.merged_into_id IS NULL ORDER BY p.full_name""",
        [18, 26, 30, 16, 30, 30, 10, 40, 60]),
    "Brokers": ("""
        SELECT broker AS "Broker", COALESCE(broker_email,'') AS "Email", sum(listings_sent) AS "Listings sent", sum(active_now) AS "Active now",
               sum(bid_on) AS "Bid on", COALESCE(to_char(max(last_interaction),'MM/DD/YY'),'') AS "Last touch"
        FROM v_broker_activity GROUP BY 1,2 ORDER BY 3 DESC""", [26, 34, 12, 10, 8, 10]),
}


def main():
    argv = sys.argv
    if "--sandbox" in argv:
        load_env(); c = psycopg.connect(os.environ["SANDBOX_DATABASE_URL"])
    else:
        c = connect()
    out = OUT
    if "--out" in argv:
        out = type(OUT)(argv[argv.index("--out") + 1])
    run_id = c.execute("INSERT INTO job_runs (job) VALUES ('tracker_export') RETURNING id").fetchone()[0]
    wb = Workbook(); wb.remove(wb.active)
    counts = {}
    for tab, (sql, widths) in TABS.items():
        cur = c.execute(sql)
        cols = [d.name for d in cur.description]
        rows = cur.fetchall()
        counts[tab] = len(rows)
        ws = wb.create_sheet(tab)
        ws["A1"] = f"{tab}  (generated from the RTP database {dt.datetime.now():%m/%d/%y %H:%M}; edit in the dashboard or tell Claude, not here)"
        ws["A1"].font = TITLE
        for j, name in enumerate(cols, 1):
            cell = ws.cell(row=3, column=j, value=name); cell.font = HDR
        for i, r in enumerate(rows, 4):
            for j, v in enumerate(r, 1):
                cell = ws.cell(row=i, column=j, value=(float(v) if hasattr(v, "quantize") else v)); cell.font = FONT
                if cols[j - 1] in ("$ invested", "Indicated $"):
                    cell.number_format = '"$"#,##0_);[Red]("$"#,##0)'
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = "A4"
    if out.exists() and "--sandbox" not in argv:
        arch = out.parent / "archive" / f"RTP_Relationship_Tracker_{dt.date.today():%Y-%m-%d}_pre-export.xlsx"
        if not arch.exists():
            shutil.copy2(out, arch)
    wb.save(out)
    c.execute("UPDATE job_runs SET finished_at=now(), status='ok', counts=%s, message=%s WHERE id=%s",
              (psycopg.types.json.Jsonb(counts), str(out), run_id))
    c.commit()
    print("wrote", out, counts)


if __name__ == "__main__":
    import psycopg.types.json  # noqa: E402
    main()
